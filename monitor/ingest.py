import hashlib
import io
import json
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from urllib.parse import urlsplit, urlunsplit
import openpyxl


def norm(value):
    value = unicodedata.normalize("NFKC", str(value or "")).casefold().replace("ё", "е")
    return " ".join(re.findall(r"[\w]+", value))


def author_key(value):
    return " ".join(sorted(norm(value).split()))


def url_key(value):
    p = urlsplit(str(value).strip())
    if p.scheme not in ("http", "https") or not p.hostname or p.username or p.password:
        raise ValueError("Нужен публичный URL http/https без учётных данных")
    return urlunsplit((p.scheme.lower(), p.netloc.lower(), p.path or "/", p.query, ""))


def stable_id(*values):
    return hashlib.sha256("\x00".join(map(str, values)).encode()).hexdigest()[:24]


def read_workbook(data, filename, store):
    if len(data) > 10_000_000:
        raise ValueError("Размер Excel превышает 10 MB")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist()) > 50_000_000:
            raise ValueError("Распакованный Excel превышает 50 MB")
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    iid = stable_id(data.hex())
    rows = []
    groups = defaultdict(list)
    for sheet in wb:
        if sheet.max_row > 50_000 or sheet.max_column > 100:
            raise ValueError("Слишком большой лист")
        headers = [str(c.value or "").strip() for c in sheet[1]]
        for number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
            if not any(c.value for c in cells):
                continue
            raw = {headers[i] or f"column_{i+1}": c.value for i, c in enumerate(cells)}
            vals = [str(c.value or "").strip() for c in cells]
            if len(vals) < 5:
                continue
            url = cells[4].hyperlink.target if cells[4].hyperlink else vals[4]
            warnings = []
            try:
                url = url_key(url)
            except ValueError as e:
                warnings.append(str(e))
            rid = stable_id(iid, sheet.title, number)
            prepared = {"title": vals[2], "author": vals[1], "format": vals[3], "rights_holder": vals[0],
                        "url": url, "domain": urlsplit(url).hostname, "title_key": norm(vals[2]),
                        "author_key": author_key(vals[1]), "warnings": warnings}
            book_id = stable_id(prepared["title_key"], prepared["author_key"], vals[3])
            prepared["book_id"] = book_id
            groups[book_id].append(rid)
            rows.append({"id": rid, "import_id": iid, "sheet": sheet.title, "row_number": number,
                         "raw": raw, "prepared": prepared})
    title_authors = defaultdict(set)
    urls = defaultdict(list)
    for row in rows:
        p = row["prepared"]
        title_authors[p["title_key"]].add(p["author_key"])
        urls[p["url"]].append(row["id"])
    for row in rows:
        p = row["prepared"]
        p["duplicate_url_rows"] = [r for r in urls[p["url"]] if r != row["id"]]
        p["same_book_rows"] = [r for r in groups[p["book_id"]] if r != row["id"]]
        if len(title_authors[p["title_key"]]) > 1:
            p["warnings"].append("Одинаковое название у разных авторов: требуется проверка каталога")
        if not p["title"] or not p["author"]:
            p["warnings"].append("Отсутствуют название или автор")
    report = {"import_id": iid, "rows": len(rows), "sheets": dict(Counter(r["sheet"] for r in rows)),
              "provisional_books": len(groups), "duplicate_url_groups": sum(len(v)>1 for v in urls.values()),
              "ambiguous_titles": sum(len(v)>1 for v in title_authors.values()),
              "catalog_kind": "provisional", "note": "Связи из исходных строк, не эталонный каталог"}
    with store.db() as db:
        db.execute("INSERT OR REPLACE INTO imports VALUES(?,?,strftime('%s','now'),?)", (iid, filename, json.dumps(report, ensure_ascii=False)))
        for row in rows:
            db.execute("INSERT OR REPLACE INTO records VALUES(?,?,?,?,?,?)", (row["id"], iid, row["sheet"], row["row_number"], json.dumps(row["raw"], ensure_ascii=False, default=str), json.dumps(row["prepared"], ensure_ascii=False)))
        for bid, ids in groups.items():
            p = next(r["prepared"] for r in rows if r["id"] == ids[0])
            existing = db.execute("SELECT source_ids,provisional FROM books WHERE id=?", (bid,)).fetchone()
            if existing and not existing["provisional"]:
                continue
            merged = sorted(set(ids + (json.loads(existing["source_ids"]) if existing else [])))
            db.execute("INSERT OR REPLACE INTO books VALUES(?,?,?,?,?,?)", (bid, p["title"], p["author"], p["format"], 1, json.dumps(merged)))
    store.event("import", f"{filename}: {len(rows)} строк; {len(groups)} предварительных книг", level="info")
    return report
